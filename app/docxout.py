import re
import docx
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from .mdparse import any_block

NAVY = RGBColor(0x0B, 0x1F, 0x4D)
BLUE = RGBColor(0x12, 0x61, 0xF5)
GREY = RGBColor(0x5F, 0x6F, 0x91)

# **dam** | *nghieng* | `code` | [chu](link)
INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")
LINK = re.compile(r"^\[([^\]]+)\]\(([^)]+)\)$")


def _hyperlink(p, text, url):
    part = p.part
    r_id = part.relate_to(url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK,
                          is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    color = OxmlElement("w:color"); color.set(qn("w:val"), "1261F5")
    u = OxmlElement("w:u"); u.set(qn("w:val"), "single")
    rPr.append(color); rPr.append(u)
    run.append(rPr)
    t = OxmlElement("w:t"); t.text = text
    run.append(t)
    link.append(run)
    p._p.append(link)


def _runs(p, text, base_italic=False, color=None):
    for part in INLINE.split(text):
        if not part:
            continue
        m = LINK.match(part)
        if m:
            _hyperlink(p, m.group(1), m.group(2))
            continue
        if part.startswith("**") and part.endswith("**"):
            r = p.add_run(part[2:-2]); r.bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            r = p.add_run(part[1:-1]); r.italic = True
        elif part.startswith("`") and part.endswith("`"):
            r = p.add_run(part[1:-1]); r.font.name = "Menlo"
        else:
            r = p.add_run(part)
        if base_italic:
            r.italic = True
        if color is not None:
            r.font.color.rgb = color


def md_to_docx(md: str, images=None, title_fallback="Nội dung"):
    """images: [{section, filename, path}] — chen anh vao cuoi muc tuong ung."""
    images = images or []
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"; st.font.size = Pt(11.5)

    lines = (md or "").splitlines()
    sec_i = 0          # dem muc theo thu tu xuat hien, khop voi mdparse
    cur_sec = None
    pending = {}
    for im in images:
        pending.setdefault(im["section"], []).append(im)

    def flush(sec):
        for im in pending.pop(sec, []):
            try:
                doc.add_picture(im["path"], width=Inches(6.3))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            except Exception:
                pass

    i = 0
    while i < len(lines):
        line = lines[i].strip()
        i += 1
        if not line:
            continue
        if line == "---":
            flush(cur_sec); continue
        if any_block(line)[0]:
            # bo nhan + mo ta minh hoa, chi giu lai anh
            while i < len(lines):
                t = lines[i].strip()
                if not t or t.startswith("#") or t == "---" or any_block(t)[0]:
                    break
                i += 1
            continue
        if line.startswith("### "):
            h = doc.add_heading(line[4:].strip(), level=2)
            for r in h.runs: r.font.color.rgb = BLUE
            continue
        if line.startswith("## "):
            flush(cur_sec)
            sec_i += 1; cur_sec = sec_i
            h = doc.add_heading(line[3:].strip(), level=1)
            for r in h.runs: r.font.color.rgb = NAVY
            continue
        if line.startswith("# "):
            h = doc.add_heading(line[2:].strip(), level=0)
            for r in h.runs: r.font.color.rgb = NAVY
            continue
        if line.startswith("> "):
            p = doc.add_paragraph(style="Intense Quote")
            _runs(p, line[2:].strip())
            continue
        if re.match(r"^[-*]\s+", line):
            p = doc.add_paragraph(style="List Bullet")
            _runs(p, re.sub(r"^[-*]\s+", "", line))
            continue
        if re.match(r"^\d+[.)]\s+", line):
            p = doc.add_paragraph(style="List Number")
            _runs(p, re.sub(r"^\d+[.)]\s+", "", line))
            continue
        p = doc.add_paragraph()
        _runs(p, line)

    flush(cur_sec)
    for k in list(pending):
        flush(k)
    return doc


def script_to_docx(text: str, title=None):
    doc = Document()
    doc.styles["Normal"].font.size = Pt(12)
    if title:
        h = doc.add_heading(title, level=0)
        for r in h.runs: r.font.color.rgb = NAVY
    for para in (text or "").split("\n"):
        if para.strip():
            doc.add_paragraph(para.strip())
    return doc
